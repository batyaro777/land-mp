"""Разбор входа брокера: текст, CSV, XLSX → список Listing.

Склейка дублей по кадастровому номеру.
"""

from __future__ import annotations

import csv
import io
import re
from collections import OrderedDict
from pathlib import Path
from typing import IO, Literal

from landcheck.models import SOTKA_M2, Listing

AreaUnit = Literal["sotka", "m2"]

CAD_RE = re.compile(r"\d{1,2}\s*:\s*\d{1,2}\s*:\s*\d{1,7}\s*:\s*\d{1,9}")
_NUM_RE = re.compile(r"[-+]?\d[\d\s ]*(?:[.,]\d+)?")
_URL_RE = re.compile(r"https?://\S+")

CAD_HEADERS = ("кадастров", "кад.", "кн", "cad", "cadastral", "cn")
PRICE_HEADERS = ("цена", "стоимост", "price", "руб")
AREA_HEADERS = ("площад", "area", "сот", "м2", "м²", "кв")
URL_HEADERS = ("ссылк", "url", "link", "объявлен")
NOTE_HEADERS = ("коммент", "примеч", "note", "описан")

# Порог: число ≥ этого считаем ценой в рублях, если колонки не подписаны.
PRICE_MIN_RUB = 50_000


def normalize_cad_num(raw: str) -> str:
    """'50 : 09 : 0060202 : 123' → '50:09:0060202:123'."""
    return re.sub(r"\s+", "", raw.strip())


def is_valid_cad_num(cn: str) -> bool:
    """Округ 2 цифры, район 2, квартал 6–7, номер участка ≥ 1."""
    return bool(re.fullmatch(r"\d{2}:\d{2}:\d{6,7}:\d{1,9}", cn))


def parse_number(token: str) -> float | None:
    """'1 200 000' → 1200000; '1,2 млн' → 1200000; '12 сот' → 12; '850 тыс' → 850000."""
    if token is None:
        return None
    t = token.strip().lower().replace(" ", " ")
    if not t:
        return None
    m = _NUM_RE.search(t)
    if not m:
        return None
    num = m.group(0).replace(" ", "").replace(" ", "")
    if "," in num and "." not in num:
        num = num.replace(",", ".")
    try:
        value = float(num)
    except ValueError:
        return None
    tail = t[m.end() :]
    if "млн" in tail or "mln" in tail:
        value *= 1_000_000
    elif "тыс" in tail or "тыc" in tail or "k" == tail.strip()[:1]:
        value *= 1_000
    return value


def area_to_m2(value: float | None, unit: AreaUnit) -> float | None:
    if value is None:
        return None
    return value * SOTKA_M2 if unit == "sotka" else value


def _unit_from_header(header: str, default: AreaUnit) -> AreaUnit:
    h = header.lower()
    if "сот" in h:
        return "sotka"
    if "м2" in h or "м²" in h or "кв" in h or "m2" in h:
        return "m2"
    return default


def _price_multiplier_from_header(header: str) -> float:
    h = header.lower()
    if "млн" in h:
        return 1_000_000
    if "тыс" in h:
        return 1_000
    return 1.0


def _match_header(name: str, keys: tuple[str, ...]) -> bool:
    n = name.strip().lower()
    return any(k in n for k in keys)


_AMOUNT_RE = re.compile(
    r"(?<![\d.,:])(\d{1,3}(?:[\s\u00a0]\d{3})+|\d+)(?:[.,](\d+))?"
    r"\s*(млн|mln|тыс|руб|₽|р\.|сот\w*|м2|м²|кв\.?\s*м|га|k\b)?",
    re.IGNORECASE,
)
_AREA_UNITS = ("сот", "м2", "м²", "кв", "га")
_PRICE_UNITS = ("млн", "mln", "тыс", "руб", "₽", "р.", "k")


def _amounts_from_text(
    text: str, area_unit: AreaUnit
) -> tuple[float | None, float | None, str | None]:
    """Числа с единицами из остатка строки → (цена ₽, площадь м², примечание)."""
    price: float | None = None
    area: float | None = None
    note: str | None = None
    unassigned: list[float] = []
    for m in _AMOUNT_RE.finditer(text):
        whole = m.group(1).replace(" ", "").replace("\u00a0", "")
        value = float(f"{whole}.{m.group(2)}" if m.group(2) else whole)
        unit = (m.group(3) or "").lower()
        if unit.startswith(_AREA_UNITS) and area is None:
            if unit.startswith("га"):
                area = value * 10_000
            elif unit.startswith("сот"):
                area = value * SOTKA_M2
            else:
                area = value
        elif unit and unit.startswith(_PRICE_UNITS) and price is None:
            if unit.startswith(("млн", "mln")):
                value *= 1_000_000
            elif unit.startswith(("тыс", "k")):
                value *= 1_000
            price = value
        else:
            unassigned.append(value)
    for value in unassigned:
        if price is None and (value >= PRICE_MIN_RUB or area is not None):
            price = value
        elif area is None:
            area = area_to_m2(value, area_unit)
            if value < PRICE_MIN_RUB and price is None:
                note = "площадь распознана эвристикой: число без подписи"
    return price, area, note


def parse_text(text: str, area_unit: AreaUnit = "sotka") -> list[Listing]:
    """Свободный текст: по строке на участок.

    Поддерживается «КН; цена; площадь» в любом разделителе, а также строки с одним
    числом: ≥ 50 000 — цена в рублях, иначе площадь. Единица площади без подписи —
    `area_unit` (по умолчанию сотки).
    """
    listings: list[Listing] = []
    for line in text.splitlines():
        if not line.strip():
            continue
        matches = list(CAD_RE.finditer(line))
        if not matches:
            continue
        # Несколько номеров в строке («КН1; КН2; КН3») — каждый со своим хвостом до следующего КН.
        for idx, m in enumerate(matches):
            seg_end = matches[idx + 1].start() if idx + 1 < len(matches) else len(line)
            rest = line[m.end() : seg_end]
            if idx == 0:
                rest = line[: m.start()] + " " + rest
            url_m = _URL_RE.search(rest)
            url = url_m.group(0) if url_m else None
            if url:
                rest = rest.replace(url, " ")
            price, area, note = _amounts_from_text(rest, area_unit)
            listings.append(
                Listing(
                    cad_num=normalize_cad_num(m.group(0)),
                    price_rub=price,
                    declared_area_m2=area,
                    url=url,
                    note=note,
                    raw_cad_num=m.group(0),
                )
            )
    return listings


def _rows_to_listings(
    rows: list[list[str]], area_unit: AreaUnit, source_name: str
) -> list[Listing]:
    if not rows:
        return []
    header = [str(c or "") for c in rows[0]]
    has_header = not any(CAD_RE.search(c) for c in header)
    if not has_header:
        # Нет заголовка — трактуем как свободный текст построчно.
        text = "\n".join(";".join(str(c or "") for c in r) for r in rows)
        return parse_text(text, area_unit)

    idx: dict[str, int] = {}
    for i, name in enumerate(header):
        if "cad" not in idx and _match_header(name, CAD_HEADERS):
            idx["cad"] = i
        elif "price" not in idx and _match_header(name, PRICE_HEADERS):
            idx["price"] = i
        elif "area" not in idx and _match_header(name, AREA_HEADERS):
            idx["area"] = i
        elif "url" not in idx and _match_header(name, URL_HEADERS):
            idx["url"] = i
        elif "note" not in idx and _match_header(name, NOTE_HEADERS):
            idx["note"] = i
    if "cad" not in idx:
        # Колонка с КН не подписана — ищем первую колонку, где встречаются КН.
        for i in range(len(header)):
            if any(i < len(r) and r[i] and CAD_RE.search(str(r[i])) for r in rows[1:]):
                idx["cad"] = i
                break
    if "cad" not in idx:
        raise ValueError(f"{source_name}: не нашёл колонку с кадастровыми номерами")

    unit = _unit_from_header(header[idx["area"]], area_unit) if "area" in idx else area_unit
    price_mult = _price_multiplier_from_header(header[idx["price"]]) if "price" in idx else 1.0

    def cell(r: list[str], key: str) -> str | None:
        i = idx.get(key)
        if i is None or i >= len(r) or r[i] is None:
            return None
        s = str(r[i]).strip()
        return s or None

    listings: list[Listing] = []
    for r in rows[1:]:
        raw_cad = cell(r, "cad")
        if not raw_cad:
            continue
        m = CAD_RE.search(raw_cad)
        if not m:
            continue
        price_raw = cell(r, "price")
        price = parse_number(price_raw) if price_raw else None
        if price is not None and price_mult != 1.0 and "млн" not in price_raw.lower():
            price *= price_mult
        area_raw = cell(r, "area")
        area = area_to_m2(parse_number(area_raw), unit) if area_raw else None
        listings.append(
            Listing(
                cad_num=normalize_cad_num(m.group(0)),
                price_rub=price,
                declared_area_m2=area,
                url=cell(r, "url"),
                note=cell(r, "note"),
                raw_cad_num=raw_cad,
            )
        )
    return listings


def parse_csv(text: str, area_unit: AreaUnit = "sotka", source_name: str = "csv") -> list[Listing]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    rows = [row for row in csv.reader(io.StringIO(text), dialect)]
    return _rows_to_listings(rows, area_unit, source_name)


def parse_xlsx(path: Path | IO[bytes], area_unit: AreaUnit = "sotka") -> list[Listing]:
    from openpyxl import load_workbook

    name = Path(path).name if isinstance(path, str | Path) else "xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        if row is None or all(c is None for c in row):
            continue
        rows.append(["" if c is None else str(c) for c in row])
    wb.close()
    return _rows_to_listings(rows, area_unit, name)


def _as_file(source: str | Path) -> Path | None:
    """Путь к существующему файлу или None. Многострочный или длинный текст — не путь:
    ``Path.exists()`` на строке длиннее 255 байт падает с OSError «File name too long»
    (пойман 14.10 на подборке из 30 номеров)."""
    if isinstance(source, Path):
        text = str(source)
    else:
        text = source
        if "\n" in text or "\r" in text or len(text.encode("utf-8", "ignore")) > 255:
            return None
    try:
        p = Path(text)
        return p if p.exists() and p.is_file() else None
    except (OSError, ValueError):
        return None


def load_listings(source: str | Path, area_unit: AreaUnit = "sotka") -> list[Listing]:
    """Файл .csv/.txt/.xlsx или строка с номерами."""
    p = _as_file(source)
    if p is not None:
        suffix = p.suffix.lower()
        if suffix in (".xlsx", ".xlsm"):
            return parse_xlsx(p, area_unit)
        text = p.read_text(encoding="utf-8-sig")
        if suffix == ".csv":
            return parse_csv(text, area_unit, p.name)
        return parse_text(text, area_unit)
    return parse_text(str(source), area_unit)


def dedupe(listings: list[Listing]) -> list[tuple[Listing, list[Listing]]]:
    """Склейка: один КН в нескольких объявлениях → (первое, [остальные])."""
    groups: OrderedDict[str, list[Listing]] = OrderedDict()
    for item in listings:
        groups.setdefault(item.cad_num, []).append(item)
    out: list[tuple[Listing, list[Listing]]] = []
    for items in groups.values():
        primary = items[0]
        # Берём первую заполненную цену/площадь среди дублей, чтобы не терять данные.
        for dup in items[1:]:
            if primary.price_rub is None and dup.price_rub is not None:
                primary.price_rub = dup.price_rub
            if primary.declared_area_m2 is None and dup.declared_area_m2 is not None:
                primary.declared_area_m2 = dup.declared_area_m2
        out.append((primary, items[1:]))
    return out
