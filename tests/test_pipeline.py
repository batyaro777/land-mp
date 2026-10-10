import json
from pathlib import Path

from landcheck.cache import ParcelCache
from landcheck.export import to_csv, to_markdown, to_xlsx
from landcheck.inputs import load_listings
from landcheck.models import Level, Purpose, Request
from landcheck.pipeline import check_batch, fetch_payload
from landcheck.providers.fixtures import FixtureProvider
from tests.helpers import make_payload


def test_check_batch_end_to_end(fixtures_dir: Path, sample_csv: Path):
    listings = load_listings(sample_csv)
    assert len(listings) == 10  # строка «не номер» не содержит КН и отбрасывается
    batch = check_batch(listings, Request(purpose=Purpose.IZHS), FixtureProvider(fixtures_dir))
    by_cn = {r.listing.cad_num: r for r in batch.results}
    assert len(batch.results) == 9  # дубль 101 склеен
    assert by_cn["50:09:0060202:101"].verdict is Level.OK
    assert len(by_cn["50:09:0060202:101"].duplicates) == 1
    assert by_cn["50:09:0060202:102"].verdict is Level.STOP  # сельхоз под ИЖС
    assert by_cn["50:09:0060202:103"].verdict is Level.STOP  # без границ
    assert by_cn["50:09:0060202:104"].verdict is Level.CHECK  # ЗОУИТ
    assert by_cn["50:09:0060202:105"].verdict is Level.STOP  # архивный
    assert by_cn["50:09:0060202:106"].verdict is Level.CHECK  # публичная собственность
    assert by_cn["50:09:0060202:107"].verdict is Level.STOP  # ВРИ магазины
    assert by_cn["50:09:0060202:999"].verdict is Level.STOP  # не найден
    # сортировка: сначала ок, потом проверить, потом стоп
    ranks = [r.verdict.rank for r in batch.results]
    assert ranks == sorted(ranks)
    assert batch.median_price_per_sotka is not None
    assert batch.counts["стоп"] == 5


def test_invalid_cad_num_is_stop_without_fetch(fixtures_dir: Path):
    listings = load_listings("1:2:3:4; 1000000; 10")
    batch = check_batch(listings, Request(), FixtureProvider(fixtures_dir))
    assert batch.results[0].verdict is Level.STOP
    assert batch.results[0].flags[0].rule_id == "R01"


def test_exports(fixtures_dir: Path, sample_csv: Path, tmp_path: Path):
    batch = check_batch(
        load_listings(sample_csv), Request(purpose=Purpose.IZHS), FixtureProvider(fixtures_dir)
    )
    xlsx = to_xlsx(batch, tmp_path / "r.xlsx")
    assert xlsx.exists() and xlsx.stat().st_size > 5000
    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Подборка", "Флаги", "ЗОУИТ", "Источники и поля", "О проверке"]
    assert wb["Подборка"].max_row == 10
    csv_text = to_csv(batch)
    assert csv_text.count("\n") == 10
    md = to_markdown(batch)
    assert "50:09:0060202:101" in md and "не является отчётом об оценке" in md


class FlakyProvider:
    name = "flaky"

    def __init__(self, payload, fail=False):
        self.payload, self.fail, self.calls = payload, fail, 0

    def fetch_raw(self, cad_num):
        self.calls += 1
        if self.fail:
            raise ConnectionError("нет сети")
        return self.payload


def test_cache_hit_and_stale_fallback(tmp_path: Path):
    cache = ParcelCache(tmp_path / "c.sqlite", ttl_days=7)
    payload = make_payload("50:09:0060202:101")
    prov = FlakyProvider(payload)
    p1, e1 = fetch_payload("50:09:0060202:101", prov, cache)
    p2, e2 = fetch_payload("50:09:0060202:101", prov, cache)
    assert prov.calls == 1 and not e1 and not e2 and p2["found"]
    # сеть упала, кэш протух → отдаём устаревший с ошибкой
    cache.ttl = cache.ttl * 0
    bad = FlakyProvider(payload, fail=True)
    p3, e3 = fetch_payload("50:09:0060202:101", bad, cache)
    assert p3["found"] and e3 and "_cache_fetched_at" in p3
    cache.close()


def test_fixture_provider_missing_file_is_not_found_with_error(tmp_path: Path):
    prov = FixtureProvider(tmp_path)
    payload = prov.fetch_raw("50:09:0060202:1")
    assert payload["found"] is False and payload["errors"]
    saved = prov.save(make_payload("50:09:0060202:2"))
    assert json.loads(saved.read_text(encoding="utf-8"))["found"] is True
